import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

D = json.load(open("data/concepts.json"))
K, L, T, M = D["KIND_NAMES"], D["LEVELS"], D["TAGS"], D["META"]
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="2F5D8A")
N = Font(name=F, size=10); LK = Font(name=F, size=10, color="2F5D8A", underline="single")
wrap = Alignment(wrap_text=True, vertical="top"); top = Alignment(vertical="top")

wb = Workbook()
# ---------------- About
ab = wb.active; ab.title = "About"
rows = [
    (M["title"], Font(name=F, size=18, bold=True, color="2F5D8A")),
    (M["edition"] + " · current to " + M["asof"], Font(name=F, size=11, color="5E666C")),
    ("", N),
    ("How to filter", Font(name=F, size=12, bold=True)),
    ("Concepts sheet: one row per concept. Use the filter arrows on the header row to filter by Category, Since (year emerged), Status, Level or any Tag column (Yes/blank).", N),
    ("Resources sheet: one row per resource link. Filter by Year, Type, Category, Level or Concept; click a URL to open it.", N),
    ("Summary sheet: live counts by category, by year of emergence and by resource type/year (formulas recalculate if you add rows).", N),
    ("", N),
    ("Column notes", Font(name=F, size=12, bold=True)),
    ("Since = year the concept emerged. Active until = latest year of major change. Status: New in 2026, Active in 2026, or Established.", N),
    ("Resource Year = publication year of an article, paper or spec, or launch year of a repository or site (approximate for long-running projects).", N),
    ("Chapter 1 (Frontier Landscape) reflects public reporting as of " + M["asof"] + " and will date quickly.", N),
]
for i, (t, f) in enumerate(rows, 1):
    c = ab.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
ab.column_dimensions["A"].width = 120

# ---------------- Concepts
cs = wb.create_sheet("Concepts")
tagkeys = [t for t in T if t not in ("new-2026", "updated-2026")]
head = ["No.", "Concept", "Alias", "Category", "Since", "Active until", "Status", "Level"] + [T[t] for t in tagkeys] + ["Definition", "Key points", "State of the art (Sept 2026)", "Resources"]
cs.append(head)
n = 0
for ci, c in enumerate(D["CHAPTERS"], 1):
    for k in c["concepts"]:
        n += 1
        status = "New in 2026" if k["since"] >= 2026 else ("Active in 2026" if k["upd"] >= 2026 else "Established")
        cs.append([n, k["name"], k["aka"], f"{ci:02d}. {c['title']}", k["since"], k["upd"], status, L[k["level"]]]
                  + ["Yes" if t in k["tags"] else "" for t in tagkeys]
                  + [k["text"], "\n".join("• " + p for p in k["points"]), k["sota"], len(k["links"])])
widths = [6, 34, 18, 34, 8, 11, 14, 13] + [12] * len(tagkeys) + [70, 70, 50, 10]
for i, w in enumerate(widths, 1): cs.column_dimensions[get_column_letter(i)].width = w
for cell in cs[1]: cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical="center")
for row in cs.iter_rows(min_row=2):
    for cell in row: cell.font = N; cell.alignment = wrap
cs.freeze_panes = "C2"
tab = Table(displayName="Concepts", ref=f"A1:{get_column_letter(len(head))}{n+1}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); cs.add_table(tab)

# ---------------- Resources
rs = wb.create_sheet("Resources")
rh = ["Year", "Type", "Resource", "URL", "Concept", "Category", "Concept since", "Level", "Status"]
rs.append(rh); r = 1
for ci, c in enumerate(D["CHAPTERS"], 1):
    for k in c["concepts"]:
        status = "New in 2026" if k["since"] >= 2026 else ("Active in 2026" if k["upd"] >= 2026 else "Established")
        for l in k["links"]:
            r += 1
            rs.append([l["year"], K[l["kind"]], l["label"], l["url"], k["name"], f"{ci:02d}. {c['title']}", k["since"], L[k["level"]], status])
            u = rs.cell(row=r, column=4); u.hyperlink = l["url"]
for i, w in enumerate([8, 10, 52, 60, 34, 34, 13, 13, 14], 1): rs.column_dimensions[get_column_letter(i)].width = w
for cell in rs[1]: cell.font = H; cell.fill = HF
for row in rs.iter_rows(min_row=2):
    for cell in row: cell.font = N; cell.alignment = top
    row[3].font = LK
rs.freeze_panes = "A2"
tab = Table(displayName="Resources", ref=f"A1:I{r}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); rs.add_table(tab)

# ---------------- Summary (formulas)
sm = wb.create_sheet("Summary")
bold = Font(name=F, bold=True)
def hdr(row, vals):
    for j, v in enumerate(vals, 1):
        c = sm.cell(row=row, column=j, value=v); c.font = H; c.fill = HF
sm["A1"] = "Live counts (recalculate automatically)"; sm["A1"].font = Font(name=F, size=13, bold=True)
row = 3; hdr(row, ["Category", "Concepts", "New in 2026", "Active in 2026", "Resources"])
cats = [f"{ci:02d}. {c['title']}" for ci, c in enumerate(D["CHAPTERS"], 1)]
start = row + 1
for cat in cats:
    row += 1
    sm.cell(row=row, column=1, value=cat)
    sm.cell(row=row, column=2, value=f'=COUNTIFS(Concepts!$D:$D,$A{row})')
    sm.cell(row=row, column=3, value=f'=COUNTIFS(Concepts!$D:$D,$A{row},Concepts!$G:$G,"New in 2026")')
    sm.cell(row=row, column=4, value=f'=COUNTIFS(Concepts!$D:$D,$A{row},Concepts!$G:$G,"<>Established")')
    sm.cell(row=row, column=5, value=f'=COUNTIFS(Resources!$F:$F,$A{row})')
row += 1; sm.cell(row=row, column=1, value="Total").font = bold
for col in "BCDE":
    c = sm[f"{col}{row}"]; c.value = f"=SUM({col}{start}:{col}{row-1})"; c.font = bold

row += 2; hdr(row, ["Year", "Concepts emerged", "Resources published"]); ys = row + 1
for y in range(2026, 2016, -1):
    row += 1
    sm.cell(row=row, column=1, value=y)
    sm.cell(row=row, column=2, value=f"=COUNTIFS(Concepts!$E:$E,$A{row})")
    sm.cell(row=row, column=3, value=f"=COUNTIFS(Resources!$A:$A,$A{row})")
row += 1; sm.cell(row=row, column=1, value="Before 2017")
sm.cell(row=row, column=2, value='=COUNTIFS(Concepts!$E:$E,"<2017")')
sm.cell(row=row, column=3, value='=COUNTIFS(Resources!$A:$A,"<2017")')

row += 2; hdr(row, ["Resource type", "All years", "2026", "2025", "2024", "Before 2024"])
for kname in K.values():
    row += 1
    sm.cell(row=row, column=1, value=kname)
    sm.cell(row=row, column=2, value=f"=COUNTIFS(Resources!$B:$B,$A{row})")
    for j, y in zip((3, 4, 5), (2026, 2025, 2024)):
        sm.cell(row=row, column=j, value=f"=COUNTIFS(Resources!$B:$B,$A{row},Resources!$A:$A,{y})")
    sm.cell(row=row, column=6, value=f'=COUNTIFS(Resources!$B:$B,$A{row},Resources!$A:$A,"<2024")')

row += 2; hdr(row, ["Level", "Concepts"])
for lv in L.values():
    row += 1; sm.cell(row=row, column=1, value=lv); sm.cell(row=row, column=2, value=f"=COUNTIFS(Concepts!$H:$H,$A{row})")
for rr in sm.iter_rows(min_row=2):
    for c in rr:
        if c.font != H and not c.font.bold: c.font = N
sm.column_dimensions["A"].width = 52
for col in "BCDEF": sm.column_dimensions[col].width = 17

wb.move_sheet("Summary", offset=-2)
wb.save("editions/AI_Concept_Book_Catalog.xlsx")
print("xlsx ok", n, r - 1)
